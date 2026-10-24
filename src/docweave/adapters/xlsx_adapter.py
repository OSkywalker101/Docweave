"""XLSX adapter (openpyxl).

Spreadsheets are the clearest case for a canonical IR: flattening a workbook into
prose destroys every relationship the sheet encodes. Formulas are kept as
formula strings and cached results are kept separately. DocWeave never evaluates
a formula, never follows an external link and never opens a macro.

Cached results are collected in a second *streaming* pass rather than by random
cell access, so a large workbook does not get parsed twice per lookup.
"""

from __future__ import annotations

import datetime

from docweave.adapters.base import Adapter, ExtractionContext, ExtractionResult
from docweave.adapters.common import ref
from docweave.exceptions import LimitExceededError, MalformedDocumentError
from docweave.models import (
    Cell,
    Confidence,
    Document,
    Locator,
    Origin,
    SourceInfo,
    Table,
    Worksheet,
)
from docweave.models.diagnostics import Category, Code

_MAX_COLUMNS = 16_384
_NO_CACHED_MARKER = " (no cached result)"


class XlsxAdapter(Adapter):
    name = "openpyxl"
    formats = ("xlsx",)
    requires = {"openpyxl": "openpyxl"}
    capabilities = (
        "worksheet identity preserved (name, index, hidden state)",
        "cell coordinates and A1 references",
        "formulas preserved verbatim alongside cached results",
        "merged ranges reported",
        "number formats retained",
        "styled-but-empty cells retained",
        "never evaluates formulas or macros",
    )
    limitations = (
        "charts, pivot tables and conditional formatting are not described",
        "cached formula results are absent if the file was never recalculated by a spreadsheet app",
        "images inside sheets are not extracted",
        "external workbook links are not followed",
    )
    cost = "moderate"

    def extract(self, ctx: ExtractionContext) -> ExtractionResult:
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            from docweave.exceptions import DependencyMissingError

            raise DependencyMissingError("openpyxl", "XLSX conversion", "office") from exc

        try:
            workbook = openpyxl.load_workbook(
                str(ctx.path), data_only=False, keep_links=False, read_only=False
            )
        except Exception as exc:
            raise MalformedDocumentError(f"XLSX could not be opened: {exc}") from exc

        sheets: list[Worksheet] = []
        pending_formulas: list[tuple[str, str, Cell]] = []
        total_formulas = 0
        hidden_sheets: list[str] = []
        truncated_sheets: list[str] = []
        defined_names: list[str] = []

        try:
            defined_names = _defined_names(workbook)
            for index, sheet in enumerate(workbook.worksheets, start=1):
                ctx.check_deadline(f"worksheet '{sheet.title}'")
                table, formula_cells, truncated = self._build_table(ctx, sheet, index)
                total_formulas += len(formula_cells)
                pending_formulas.extend((sheet.title, coord, cell) for coord, cell in formula_cells)

                if sheet.sheet_state != "visible":
                    hidden_sheets.append(sheet.title)
                    ctx.warn(
                        Code.HIDDEN_CONTENT,
                        f"Worksheet '{sheet.title}' is '{sheet.sheet_state}'; its content was converted.",
                        category=Category.UNVERIFIED,
                        severity="info",
                        reason="Hidden sheets are included by default; most viewers hide them.",
                        scope_kind="sheet",
                        scope_name=sheet.title,
                        evidence={"sheet_state": sheet.sheet_state},
                    )

                merged = len(sheet.merged_cells.ranges) if sheet.merged_cells is not None else 0
                if merged:
                    ctx.warn(
                        Code.EMPTY_CELLS,
                        f"Worksheet '{sheet.title}' has {merged} merged range(s); only the "
                        "anchor cell carries a value.",
                        category=Category.HEURISTIC,
                        severity="info",
                        reason="Merged ranges are reported rather than duplicated across cells.",
                        scope_kind="sheet",
                        scope_name=sheet.title,
                        evidence={"merged_ranges": merged},
                    )

                if truncated:
                    truncated_sheets.append(sheet.title)
                    ctx.warn(
                        Code.LIMIT_EXCEEDED,
                        f"Worksheet '{sheet.title}' was truncated to {table.meta.get('converted_rows')} rows.",
                        category=Category.INTERVENTION,
                        reason="Per-sheet row budget reached.",
                        hint="Raise ConversionConfig.xlsx_max_rows_per_sheet to include more.",
                        scope_kind="sheet",
                        scope_name=sheet.title,
                    )

                sheets.append(
                    Worksheet(
                        name=sheet.title,
                        index=index,
                        hidden=sheet.sheet_state != "visible",
                        dimension=_dimension(sheet),
                        table=table,
                        formula_count=len(formula_cells),
                        merged_range_count=merged,
                        truncated=truncated,
                    )
                )
        finally:
            workbook.close()

        missing_cached = self._attach_cached_values(ctx, pending_formulas)

        if missing_cached:
            ctx.warn(
                Code.FORMULA_WITHOUT_CACHED_VALUE,
                f"{missing_cached} formula cell(s) have no cached result. DocWeave does not "
                "evaluate formulas, so their computed values are reported as unavailable.",
                category=Category.UNVERIFIED,
                reason="Cached values exist only after a spreadsheet application recalculates the file.",
                hint="Open and save the workbook in Excel or LibreOffice to populate cached results.",
                evidence={"formula_cells_without_cached_value": missing_cached},
            )

        document = Document(
            adapter=self.name,
            container_kind="sheets",
            sheets=sheets,
            metadata={
                "sheet_count": len(sheets),
                "formula_count": total_formulas,
                "hidden_sheets": hidden_sheets,
                "truncated_sheets": truncated_sheets,
                "defined_names": defined_names,
            },
        )
        return ExtractionResult(
            document=document,
            source=SourceInfo(name=ctx.path.name, byte_size=ctx.path.stat().st_size),
            container_kind="sheets",
        )

    # -- pass 1: structure and formulas ------------------------------------

    def _build_table(
        self, ctx: ExtractionContext, sheet, sheet_index: int
    ) -> tuple[Table, list[tuple[str, Cell]], bool]:
        max_row = sheet.max_row or 0
        max_col = min(sheet.max_column or 0, _MAX_COLUMNS)
        limit = ctx.config.xlsx_max_rows_per_sheet
        converted_rows = min(max_row, limit)
        truncated = max_row > converted_rows

        rows: list[list[Cell]] = []
        formula_cells: list[tuple[str, Cell]] = []
        cell_budget = ctx.config.max_cells

        for r in range(1, converted_rows + 1):
            cells: list[Cell] = []
            for c in range(1, max_col + 1):
                cell = sheet.cell(row=r, column=c)
                value = cell.value
                if value is None and not _has_style(cell):
                    continue
                cell_budget -= 1
                if cell_budget < 0:
                    raise LimitExceededError(
                        f"Worksheet '{sheet.title}' exceeds the cell budget of "
                        f"{ctx.config.max_cells}. Narrow the sheet range, or raise "
                        "ConversionConfig.max_cells."
                    )
                locator = Locator(
                    kind="sheet", sheet=sheet.title, cell=cell.coordinate, row=r, column=c
                )
                is_formula = isinstance(value, str) and value.startswith("=")
                text = "" if value is None else _cell_display(value)
                new_cell = Cell(
                    id=ctx.ids.child(None, "c"),
                    text=text,
                    row=r - 1,
                    column=c - 1,
                    is_header=(r == 1),
                    formula=value if is_formula else None,
                    value=None if is_formula else text,
                    number_format=cell.number_format,
                    locator=locator,
                )
                if is_formula:
                    formula_cells.append((cell.coordinate, new_cell))
                cells.append(new_cell)
            rows.append(cells)

        while rows and not rows[-1]:
            rows.pop()

        table = Table(
            id=ctx.ids.child(None, "table"),
            text="",
            rows=rows,
            header_rows=0,
            structure="grid",
            sparse=True,
            sources=[
                ref(
                    self.name,
                    Locator(kind="sheet", sheet=sheet.title),
                    origin=Origin.OBSERVED,
                    confidence=Confidence.HIGH,
                    note=(
                        "Worksheet grid. Cells with no value and no styling are omitted, but "
                        "every retained cell keeps its A1 reference, row and column, so the "
                        "sheet geometry is fully recoverable from the JSON output."
                    ),
                )
            ],
            meta={"sheet": sheet.title, "row_count": len(rows), "converted_rows": converted_rows},
        )
        del sheet_index
        return table, formula_cells, truncated

    # -- pass 2: cached formula results ------------------------------------

    def _attach_cached_values(
        self, ctx: ExtractionContext, pending: list[tuple[str, str, Cell]]
    ) -> int:
        """Stream the workbook a second time to pick up cached formula results.

        Returns the number of formula cells with no cached value, which the
        caller reports as unverified content.
        """
        if not ctx.config.xlsx_cached_values or not pending:
            return 0

        try:
            import openpyxl  # type: ignore
        except ImportError:  # pragma: no cover - declared dependency
            return 0

        wanted: dict[str, set[str]] = {}
        for sheet_title, coord, _ in pending:
            wanted.setdefault(sheet_title, set()).add(coord)

        resolved: dict[tuple[str, str], object] = {}
        load_failed: str | None = None
        book = None
        try:
            book = openpyxl.load_workbook(
                str(ctx.path), data_only=True, keep_links=False, read_only=True
            )
        except Exception as exc:
            load_failed = str(exc)

        if book is not None:
            try:
                for sheet_title, coords in wanted.items():
                    try:
                        worksheet = book[sheet_title]
                    except Exception:
                        continue
                    try:
                        for row in worksheet.iter_rows():
                            for cell in row:
                                coord = cell.coordinate
                                if coord in coords and cell.value is not None:
                                    resolved[(sheet_title, coord)] = cell.value
                    except Exception:
                        continue
            finally:
                book.close()

        missing = 0
        for sheet_title, coord, cell in pending:
            if cell.formula is None:
                continue
            cached = resolved.get((sheet_title, coord))
            if cached is None:
                cell.text = f"{cell.formula}  {_NO_CACHED_MARKER}"
                cell.warnings.append("formula-without-cached-value")
                missing += 1
            else:
                cell.text = f"{cell.formula} -> {cached}"
                cell.value = str(cached)

        if load_failed is not None:
            ctx.warn(
                Code.FORMULA_WITHOUT_CACHED_VALUE,
                f"Cached formula values could not be read ({load_failed}); "
                "formulas are reported without a result.",
                category=Category.UNVERIFIED,
                reason="The second, data_only workbook load failed.",
            )
        return missing


def _dimension(sheet) -> str | None:
    try:
        if not sheet.max_row or not sheet.max_column:
            return None
        return f"A1:{sheet.cell(row=sheet.max_row, column=sheet.max_column).coordinate}"
    except Exception:
        return None


def _has_style(cell) -> bool:
    try:
        return bool(cell.has_style)
    except Exception:
        return False


def _cell_display(value) -> str:
    """Render a cell value without inventing precision it does not have.

    Numbers keep their raw value rather than the workbook's display string:
    ``0.075`` stays ``0.075`` and not ``7.5%``, because a consumer that wants a
    percentage can format it and one that wants the number cannot un-format it.

    Dates are the exception. ``str(datetime)`` on a whole-day cell emits
    ``2026-03-14 00:00:00``, and the trailing midnight is noise the reader has
    to strip -- so a zero clock time is dropped.
    """
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, datetime.time):
        return value.isoformat()
    return str(value)


def _defined_names(workbook) -> list[str]:
    try:
        return sorted(str(n) for n in workbook.defined_names)
    except Exception:
        return []
